"""Dosya analizi — PDF / DOCX / XLSX / CSV / TXT oku, gerçek AI ile özetle."""

import io
from datetime import datetime, timezone

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile

from lib.activity import log
from lib.auth import require_session
from lib.db import db
from lib.llm import LLMUnavailable, complete, lang_rule
from models.schemas import FileDoc

router = APIRouter(prefix="/files", tags=["files"], dependencies=[Depends(require_session)])

MAX_BYTES = 12 * 1024 * 1024
MAX_CHARS = 60_000


def _extract(name: str, data: bytes) -> tuple[str, str]:
    """Returns (kind, text). Raises ValueError for unsupported types."""
    lower = name.lower()
    if lower.endswith(".pdf"):
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        pages = [(p.extract_text() or "") for p in reader.pages[:80]]
        return "pdf", "\n\n".join(pages)
    if lower.endswith(".docx"):
        import docx

        doc = docx.Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return "docx", "\n".join(parts)
    if lower.endswith((".xlsx", ".xlsm")):
        import pandas as pd

        sheets = pd.read_excel(io.BytesIO(data), sheet_name=None)
        out = []
        for sheet, frame in sheets.items():
            out.append(f"### Sayfa: {sheet} ({len(frame)} satır, {len(frame.columns)} kolon)")
            out.append(frame.head(200).to_csv(index=False))
        return "xlsx", "\n".join(out)
    if lower.endswith(".csv"):
        import pandas as pd

        frame = pd.read_csv(io.BytesIO(data))
        return "csv", f"({len(frame)} satır)\n" + frame.head(300).to_csv(index=False)
    if lower.endswith((".txt", ".md", ".json", ".log")):
        return "text", data.decode("utf-8", errors="replace")
    raise ValueError("Desteklenmeyen dosya türü. PDF, DOCX, XLSX, CSV veya TXT yükleyin.")


@router.get("", response_model=list[FileDoc])
async def list_files():
    docs = await db.files.find().sort("created_at", -1).to_list(200)
    return [FileDoc(**d) for d in docs]


@router.delete("/{file_id}")
async def delete_file(file_id: str):
    res = await db.files.delete_one({"id": file_id})
    if not res.deleted_count:
        raise HTTPException(status_code=404, detail="Dosya bulunamadı.")
    await log("dosya", "Dosya kaydı silindi.")
    return {"deleted": True}


@router.get("/{file_id}", response_model=FileDoc)
async def get_file(file_id: str):
    doc = await db.files.find_one({"id": file_id})
    if not doc:
        raise HTTPException(status_code=404, detail="Dosya bulunamadı.")
    return FileDoc(**doc)


@router.post("/analyze", response_model=FileDoc)
async def analyze(file: UploadFile = File(...), lang: str = "auto"):
    data = await file.read()
    if not data:
        raise HTTPException(status_code=422, detail="Dosya boş.")
    if len(data) > MAX_BYTES:
        raise HTTPException(status_code=413, detail="Dosya 12 MB sınırını aşıyor.")

    try:
        kind, text = _extract(file.filename or "dosya", data)
    except ValueError as exc:
        raise HTTPException(status_code=415, detail=str(exc)) from exc
    except Exception as exc:  # noqa: BLE001 — corrupt/locked file, say so plainly
        await log("dosya", f"Dosya okunamadı: {file.filename} ({exc})", status="error")
        raise HTTPException(status_code=422, detail=f"Dosya okunamadı: {exc}") from exc

    text = text.strip()
    if not text:
        raise HTTPException(
            status_code=422,
            detail="Dosyadan metin çıkarılamadı (taranmış görsel PDF olabilir; OCR bu sürümde bağlı değil).",
        )
    text = text[:MAX_CHARS]

    system = (
        "Sen LEO'sun — LeoHoca'nın dosya analisti. Verilen belgeyi analiz edip yapılandırılmış "
        "özet çıkarırsın. Belgede olmayan bilgiyi asla uydurmazsın; eksikse 'belgede yok' dersin.\n"
        + lang_rule(lang)
    )
    prompt = (
        f"Dosya: {file.filename} (tür: {kind})\n\n"
        "Şu başlıklarla özetle:\n"
        "1) ÖZET (3-5 madde)\n2) ÖNEMLİ SAYILAR/TARİHLER\n3) TARAFLAR/KİŞİLER (belgede geçenler)\n"
        "4) RİSKLER VEYA DİKKAT EDİLECEKLER\n5) ÖNERİLEN AKSİYONLAR\n\n"
        f"--- BELGE İÇERİĞİ ---\n{text}"
    )
    try:
        summary = await complete(f"file-{file.filename}", system, prompt)
    except LLMUnavailable as exc:
        await log("dosya", f"Dosya analizi yapılamadı: {exc}", status="error")
        raise HTTPException(status_code=503, detail=str(exc)) from exc

    doc = FileDoc(
        filename=file.filename or "dosya",
        kind=kind,
        size_bytes=len(data),
        chars=len(text),
        summary=summary.strip(),
        excerpt=text[:4000],
        created_at=datetime.now(timezone.utc),
    )
    await db.files.insert_one(doc.model_dump())
    await log("dosya", f"Dosya analiz edildi: {doc.filename} ({kind}, {len(text)} karakter)")
    return doc
